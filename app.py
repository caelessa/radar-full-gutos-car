from __future__ import annotations

import os
import sqlite3
from datetime import datetime
from zoneinfo import ZoneInfo
from pathlib import Path
from typing import Any
from decimal import Decimal, InvalidOperation
from urllib.parse import urlparse, urlunparse
import unicodedata

import pandas as pd
import psycopg
from psycopg.rows import dict_row
from flask import Flask, flash, redirect, render_template, request, send_file, url_for

BASE_DIR = Path(__file__).resolve().parent
SEED_SQLITE_PATH = BASE_DIR / "radar_full_gutos.db"
UPLOAD_DIR = BASE_DIR / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "radar-full-gutos-car")

FUSO_BRASIL = ZoneInfo("America/Sao_Paulo")

def agora_brasil() -> datetime:
    return datetime.now(FUSO_BRASIL)

def agora_brasil_str() -> str:
    return agora_brasil().strftime("%d/%m/%Y %H:%M:%S")


def normalize_database_url(url: str) -> str:
    """Aceita postgres:// e postgresql://."""
    if url.startswith("postgres://"):
        return "postgresql://" + url[len("postgres://"):]
    return url


def get_conn():
    database_url = os.environ.get("DATABASE_URL")
    if not database_url:
        raise RuntimeError(
            "DATABASE_URL não configurada. Crie um PostgreSQL e cadastre a variável DATABASE_URL no Render."
        )
    return psycopg.connect(normalize_database_url(database_url), row_factory=dict_row)



ANUNCIOS_FULL_SCHEMA_COLUMNS = [
    ("id", "SERIAL PRIMARY KEY"),
    ("codigo_anuncio", "TEXT NOT NULL UNIQUE"),
    ("numero_produto", "TEXT"),
    ("titulo", "TEXT"),
    ("variacoes", "TEXT"),
    ("quantidade_full", "INTEGER NOT NULL DEFAULT 0"),
    ("preco", "NUMERIC(12,2) NOT NULL DEFAULT 0"),
    ("moeda", "TEXT"),
    ("condicao", "TEXT"),
    ("forma_entrega", "TEXT"),
    ("tipo_anuncio", "TEXT"),
    ("status", "TEXT"),
    ("altura_cm", "INTEGER DEFAULT 0"),
    ("largura_cm", "INTEGER DEFAULT 0"),
    ("profundidade_cm", "INTEGER DEFAULT 0"),
    ("peso_kg", "NUMERIC(12,3) DEFAULT 0"),
    ("estoque_minimo", "INTEGER NOT NULL DEFAULT 0"),
    ("estoque_recomendado", "INTEGER NOT NULL DEFAULT 0"),
    ("quantidade_enviar_full", "INTEGER GENERATED ALWAYS AS (GREATEST((CASE WHEN estoque_recomendado > 0 THEN estoque_recomendado ELSE estoque_minimo END) - quantidade_full, 0)) STORED"),
    ("precisa_repor", "TEXT GENERATED ALWAYS AS (CASE WHEN estoque_minimo > 0 AND quantidade_full <= estoque_minimo THEN 'SIM' ELSE 'NAO' END) STORED"),
    ("criado_em", "TIMESTAMP DEFAULT CURRENT_TIMESTAMP"),
    ("atualizado_em", "TIMESTAMP DEFAULT CURRENT_TIMESTAMP"),
    ("ativo_no_relatorio", "TEXT DEFAULT 'SIM'"),
    ("data_ultima_importacao", "TEXT"),
    ("data_primeira_importacao", "TEXT"),
    ("observacao", "TEXT"),
    ("custo_produto", "NUMERIC(12,2) NOT NULL DEFAULT 0"),
    ("custo_mercado_livre", "NUMERIC(12,2) NOT NULL DEFAULT 0"),
    ("custo_impostos", "NUMERIC(12,2) NOT NULL DEFAULT 0"),
]

BASE_COPY_COLUMNS = [
    "codigo_anuncio", "numero_produto", "titulo", "variacoes", "quantidade_full", "preco", "moeda", "condicao",
    "forma_entrega", "tipo_anuncio", "status", "altura_cm", "largura_cm", "profundidade_cm", "peso_kg",
    "estoque_minimo", "estoque_recomendado", "criado_em", "atualizado_em", "ativo_no_relatorio",
    "data_ultima_importacao", "data_primeira_importacao", "observacao", "custo_produto", "custo_mercado_livre", "custo_impostos"
]


def get_existing_columns(cur, table_name: str = "anuncios_full") -> set[str]:
    cur.execute(
        """
        SELECT column_name
        FROM information_schema.columns
        WHERE table_schema = 'public' AND table_name = %s
        """,
        (table_name,),
    )
    return {row["column_name"] for row in cur.fetchall()}


def create_anuncios_full_table(cur, table_name: str = "anuncios_full") -> None:
    cols_sql = ",\n                    ".join(f"{name} {definition}" for name, definition in ANUNCIOS_FULL_SCHEMA_COLUMNS)
    cur.execute(f"CREATE TABLE IF NOT EXISTS {table_name} (\n                    {cols_sql}\n                )")


def rebuild_anuncios_full_table(cur) -> None:
    """Reconstrói anuncios_full para limpar metadados internos acumulados no PostgreSQL.

    A versão antiga fez muitos DROP/ADD COLUMN de colunas geradas, e o PostgreSQL mantém
    colunas removidas internamente. Quando o limite interno de 1600 colunas é atingido,
    qualquer ALTER TABLE falha. Esta função cria uma tabela limpa, copia os dados úteis,
    renomeia a antiga como backup e coloca a nova no lugar.
    """
    existing = get_existing_columns(cur, "anuncios_full")
    backup_name = "anuncios_full_backup_" + datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    new_name = "anuncios_full_reparada"

    cur.execute(f"DROP TABLE IF EXISTS {new_name}")
    create_anuncios_full_table(cur, new_name)

    insert_cols = [c for c in BASE_COPY_COLUMNS if c in [name for name, _ in ANUNCIOS_FULL_SCHEMA_COLUMNS]]
    select_exprs = []
    for col in insert_cols:
        if col in existing:
            select_exprs.append(col)
        else:
            if col in {"quantidade_full", "altura_cm", "largura_cm", "profundidade_cm", "estoque_minimo", "estoque_recomendado"}:
                select_exprs.append(f"0 AS {col}")
            elif col in {"preco", "peso_kg", "custo_produto", "custo_mercado_livre", "custo_impostos"}:
                select_exprs.append(f"0 AS {col}")
            elif col == "ativo_no_relatorio":
                select_exprs.append("'SIM' AS ativo_no_relatorio")
            elif col in {"criado_em", "atualizado_em"}:
                select_exprs.append(f"CURRENT_TIMESTAMP AS {col}")
            else:
                select_exprs.append(f"NULL AS {col}")

    # codigo_anuncio é obrigatório; só copia linhas válidas.
    cur.execute(
        f"""
        INSERT INTO {new_name} ({', '.join(insert_cols)})
        SELECT {', '.join(select_exprs)}
        FROM anuncios_full
        WHERE codigo_anuncio IS NOT NULL AND codigo_anuncio <> ''
        ON CONFLICT (codigo_anuncio) DO NOTHING
        """
    )

    cur.execute(f"ALTER TABLE anuncios_full RENAME TO {backup_name}")
    cur.execute(f"ALTER TABLE {new_name} RENAME TO anuncios_full")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_anuncios_full_ativo ON anuncios_full(ativo_no_relatorio)")


def ensure_anuncios_columns(cur) -> None:
    required = {"observacao", "custo_produto", "custo_mercado_livre", "custo_impostos"}
    existing = get_existing_columns(cur, "anuncios_full")
    missing = required - existing
    if not missing:
        return

    # Se faltam colunas novas, preferimos reconstruir a tabela em vez de executar ALTER TABLE.
    # Isso evita o erro TooManyColumns em bancos que já acumularam metadados internos.
    rebuild_anuncios_full_table(cur)


def init_db() -> None:
    """Cria a tabela no PostgreSQL e carrega o banco seed SQLite se estiver vazio."""
    with get_conn() as conn:
        with conn.cursor() as cur:
            create_anuncios_full_table(cur, "anuncios_full")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_anuncios_full_ativo ON anuncios_full(ativo_no_relatorio)")
            ensure_anuncios_columns(cur)
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS importacoes_relatorios (
                    id SERIAL PRIMARY KEY,
                    data_importacao TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    data_importacao_br TEXT,
                    nome_arquivo TEXT,
                    total_relatorio INTEGER DEFAULT 0,
                    qtd_novos INTEGER DEFAULT 0,
                    qtd_alterados INTEGER DEFAULT 0,
                    qtd_ausentes INTEGER DEFAULT 0,
                    criado_em TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
                """
            )
            # IMPORTANTE:
            # Não fazemos DROP/ADD das colunas geradas a cada abertura da página.
            # A versão anterior repetia ALTER TABLE em todo acesso e o PostgreSQL
            # acumulava colunas internas removidas, causando:
            # psycopg.errors.TooManyColumns: tables can have at most 1600 columns.
            #
            # A regra de envio/reposição agora é calculada nas consultas principais
            # quando necessário, e as colunas geradas existentes são mantidas.

            cur.execute("SELECT COUNT(*) AS total FROM anuncios_full")
            total = cur.fetchone()["total"]
        conn.commit()

    if total == 0 and SEED_SQLITE_PATH.exists():
        seed_from_sqlite()


def seed_from_sqlite() -> None:
    """Importa os dados iniciais do radar_full_gutos.db para o PostgreSQL uma única vez."""
    sqlite_conn = sqlite3.connect(SEED_SQLITE_PATH)
    sqlite_conn.row_factory = sqlite3.Row
    rows = sqlite_conn.execute("SELECT * FROM anuncios_full").fetchall()
    sqlite_conn.close()

    if not rows:
        return

    with get_conn() as conn:
        with conn.cursor() as cur:
            for row in rows:
                r = dict(row)
                cur.execute(
                    """
                    INSERT INTO anuncios_full (
                        codigo_anuncio, numero_produto, titulo, variacoes, quantidade_full, preco, moeda, condicao,
                        forma_entrega, tipo_anuncio, status, altura_cm, largura_cm, profundidade_cm, peso_kg,
                        estoque_minimo, estoque_recomendado, ativo_no_relatorio,
                        data_primeira_importacao, data_ultima_importacao, observacao,
                        custo_produto, custo_mercado_livre, custo_impostos
                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                    ON CONFLICT (codigo_anuncio) DO NOTHING
                    """,
                    (
                        r.get("codigo_anuncio"), r.get("numero_produto"), r.get("titulo"), r.get("variacoes"),
                        int(r.get("quantidade_full") or 0), float(r.get("preco") or 0), r.get("moeda"),
                        r.get("condicao"), r.get("forma_entrega"), r.get("tipo_anuncio"), r.get("status"),
                        int(r.get("altura_cm") or 0), int(r.get("largura_cm") or 0), int(r.get("profundidade_cm") or 0),
                        float(r.get("peso_kg") or 0), int(r.get("estoque_minimo") or 0),
                        int(r.get("estoque_recomendado") or 0), r.get("ativo_no_relatorio") or "SIM",
                        r.get("data_primeira_importacao"), r.get("data_ultima_importacao"), r.get("observacao"),
                        float(r.get("custo_produto") or 0), float(r.get("custo_mercado_livre") or 0),
                        float(r.get("custo_impostos") or 0),
                    ),
                )
        conn.commit()


def int_safe(value: Any, default: int = 0) -> int:
    if pd.isna(value):
        return default
    try:
        text = str(value).strip()
        if text in {"", "-"}:
            return default
        return int(float(text.replace(",", ".")))
    except Exception:
        return default


def float_safe(value: Any, default: float = 0.0) -> float:
    if pd.isna(value):
        return default
    if isinstance(value, (int, float)):
        return float(value)
    try:
        text = str(value).strip()
        if text in {"", "-"}:
            return default
        text = text.replace("R$", "").strip()
        if "," in text:
            text = text.replace(".", "").replace(",", ".")
        return float(text)
    except Exception:
        return default


def get_series(df: pd.DataFrame, column: str, default: Any = "") -> pd.Series:
    if column in df.columns:
        return df[column]
    return pd.Series([default] * len(df), index=df.index)


EDITABLE_COLUMNS = [
    "estoque_minimo",
    "estoque_recomendado",
    "observacao",
    "custo_produto",
    "custo_mercado_livre",
    "custo_impostos",
]

EDITABLE_ALIASES = {
    "estoque_minimo": ["ESTOQUE_MINIMO", "ESTOQUE MINIMO", "ESTOQUE MÍNIMO", "MINIMO", "MÍNIMO"],
    "estoque_recomendado": ["ESTOQUE_RECOMENDADO", "ESTOQUE RECOMENDADO", "RECOMENDADO"],
    "observacao": ["OBSERVACAO", "OBSERVAÇÃO", "OBSERVACOES", "OBSERVAÇÕES", "OBS"],
    "custo_produto": ["CUSTO_PRODUTO", "CUSTO PRODUTO", "CUSTO DO PRODUTO"],
    "custo_mercado_livre": ["CUSTO_MERCADO_LIVRE", "CUSTO MERCADO LIVRE", "CUSTO ML", "TAXA ML", "TAXA MERCADO LIVRE"],
    "custo_impostos": ["CUSTO_IMPOSTOS", "CUSTO IMPOSTOS", "IMPOSTOS", "CUSTO IMPOSTO"],
}


def normalize_header(value: Any) -> str:
    text = str(value or "").strip().upper()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    for ch in ["_", "-", ".", "/", "\\", "\n", "\r", "\t"]:
        text = text.replace(ch, " ")
    return " ".join(text.split())


def find_column(df: pd.DataFrame, aliases: list[str]) -> str | None:
    normalized = {normalize_header(c): c for c in df.columns}
    for alias in aliases:
        key = normalize_header(alias)
        if key in normalized:
            return normalized[key]
    return None


def is_blank(value: Any) -> bool:
    if pd.isna(value):
        return True
    return str(value).strip() == ""


def optional_int(value: Any) -> int | None:
    if is_blank(value):
        return None
    return int_safe(value, 0)


def optional_float(value: Any) -> float | None:
    if is_blank(value):
        return None
    return float_safe(value, 0.0)


def optional_text(value: Any) -> str | None:
    if is_blank(value):
        return None
    return str(value).strip()


def extract_editable_values(df: pd.DataFrame, row_index: Any) -> dict[str, Any]:
    """Lê colunas editáveis opcionais de uma linha de Excel.

    Regra: célula em branco mantém o valor existente no banco.
    Se a coluna não existir, também mantém o valor existente.
    """
    values: dict[str, Any] = {}
    for canonical in EDITABLE_COLUMNS:
        col = find_column(df, EDITABLE_ALIASES[canonical])
        if not col:
            values[canonical] = None
            continue
        raw = df.at[row_index, col]
        if canonical in {"estoque_minimo", "estoque_recomendado"}:
            values[canonical] = optional_int(raw)
        elif canonical in {"custo_produto", "custo_mercado_livre", "custo_impostos"}:
            values[canonical] = optional_float(raw)
        else:
            values[canonical] = optional_text(raw)
    return values


def merge_editable(old: dict | None, imported: dict[str, Any]) -> dict[str, Any]:
    merged = {}
    for col in EDITABLE_COLUMNS:
        default = "" if col == "observacao" else 0
        current = old.get(col, default) if old else default
        merged[col] = imported[col] if imported.get(col) is not None else current
        if col != "observacao":
            merged[col] = float(merged[col] or 0) if col.startswith("custo_") else int(merged[col] or 0)
        else:
            merged[col] = merged[col] or ""
    return merged


def normalize_for_compare(value: Any, col: str) -> Any:
    """Normaliza valores antes de comparar importações.

    Evita falsos positivos em campos numéricos vindos do PostgreSQL como Decimal
    e do Excel como int/float, por exemplo 1.200 vs 1.2 ou 0.000 vs 0.
    """
    if value is None:
        value = 0 if col in NUMERIC_COMPARE_COLS else ""

    if col in INTEGER_COMPARE_COLS:
        return int_safe(value, 0)

    if col in FLOAT_COMPARE_COLS:
        try:
            if isinstance(value, Decimal):
                return round(float(value), 3)
            return round(float_safe(value, 0.0), 3)
        except Exception:
            return 0.0

    return str(value).strip()


INTEGER_COMPARE_COLS = {"quantidade_full", "altura_cm", "largura_cm", "profundidade_cm", "estoque_minimo", "estoque_recomendado"}
FLOAT_COMPARE_COLS = {"preco", "peso_kg", "custo_produto", "custo_mercado_livre", "custo_impostos"}
NUMERIC_COMPARE_COLS = INTEGER_COMPARE_COLS | FLOAT_COMPARE_COLS

def read_report(path: Path) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name="Anúncios", header=0, skiprows=[1, 2, 3, 4])
    df = df.dropna(how="all")

    required = ["ITEM_ID", "TITLE", "QUANTITY"]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(f"Colunas obrigatórias não encontradas: {missing}")

    out = pd.DataFrame()
    out["codigo_anuncio"] = df["ITEM_ID"].fillna("").astype(str).str.strip()
    out["numero_produto"] = get_series(df, "PRODUCT_NUMBER", "").fillna("").astype(str).str.strip()
    out["titulo"] = df["TITLE"].fillna("").astype(str).str.strip()
    out["variacoes"] = get_series(df, "VARIATIONS", "").fillna("").astype(str).str.strip()
    out["quantidade_full"] = df["QUANTITY"].apply(int_safe)
    out["preco"] = get_series(df, "PRICE", 0).apply(float_safe)
    out["moeda"] = get_series(df, "CURRENCY_ID", "").fillna("").astype(str).str.strip()
    out["condicao"] = get_series(df, "CONDITION", "").fillna("").astype(str).str.strip()
    out["forma_entrega"] = get_series(df, "SHIPPING_METHOD", "").fillna("").astype(str).str.strip()
    out["tipo_anuncio"] = get_series(df, "LISTING_TYPE", "").fillna("").astype(str).str.strip()
    out["status"] = get_series(df, "STATUS", "").fillna("").astype(str).str.strip()
    out["altura_cm"] = get_series(df, "SHIPPING_HEIGHT", 0).apply(int_safe)
    out["largura_cm"] = get_series(df, "SHIPPING_WIDTH", 0).apply(int_safe)
    out["profundidade_cm"] = get_series(df, "SHIPPING_DEPTH", 0).apply(int_safe)
    out["peso_kg"] = get_series(df, "SHIPPING_WEIGHT", 0).apply(float_safe)

    # Colunas editáveis opcionais que podem ser digitadas na própria planilha de anúncios.
    for canonical in EDITABLE_COLUMNS:
        values = []
        col = find_column(df, EDITABLE_ALIASES[canonical])
        for idx in df.index:
            if not col:
                values.append(None)
                continue
            raw = df.at[idx, col]
            if canonical in {"estoque_minimo", "estoque_recomendado"}:
                values.append(optional_int(raw))
            elif canonical in {"custo_produto", "custo_mercado_livre", "custo_impostos"}:
                values.append(optional_float(raw))
            else:
                values.append(optional_text(raw))
        out[canonical] = values

    out = out[out["codigo_anuncio"].str.startswith("MLB")].copy()
    out = out.drop_duplicates(subset=["codigo_anuncio"], keep="last")
    return out


@app.route("/")
def index():
    init_db()
    busca = request.args.get("busca", "").strip()
    filtro = request.args.get("filtro", "ativos")

    where = []
    params: list[Any] = []

    if busca:
        where.append("(codigo_anuncio ILIKE %s OR numero_produto ILIKE %s OR titulo ILIKE %s)")
        termo = f"%{busca}%"
        params.extend([termo, termo, termo])

    if filtro == "ativos":
        where.append("ativo_no_relatorio = 'SIM'")
    elif filtro == "inativos":
        where.append("ativo_no_relatorio = 'NAO'")
    elif filtro == "status_ativo":
        where.append("ativo_no_relatorio = 'SIM' AND status ILIKE %s")
        params.append('ativo%')
    elif filtro == "status_inativo":
        where.append("ativo_no_relatorio = 'SIM' AND status ILIKE %s")
        params.append('inativo%')
    elif filtro == "repor":
        where.append("ativo_no_relatorio = 'SIM' AND precisa_repor = 'SIM'")
    elif filtro == "enviar":
        where.append("ativo_no_relatorio = 'SIM' AND quantidade_enviar_full > 0")
    elif filtro in {"sem_config", "novos"}:
        where.append("ativo_no_relatorio = 'SIM' AND estoque_minimo = 0 AND estoque_recomendado = 0")

    sql = "SELECT * FROM anuncios_full"
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY ativo_no_relatorio DESC, quantidade_enviar_full DESC, titulo ASC"

    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(sql, params)
            rows = cur.fetchall()
            cur.execute(
                """
                SELECT
                    COUNT(*) AS total_base,
                    COALESCE(SUM(CASE WHEN ativo_no_relatorio = 'SIM' THEN 1 ELSE 0 END), 0) AS total_anuncios,
                    COALESCE(SUM(CASE WHEN ativo_no_relatorio = 'NAO' THEN 1 ELSE 0 END), 0) AS total_inativos,
                    COALESCE(SUM(CASE WHEN ativo_no_relatorio = 'SIM' AND status ILIKE 'inativo%%' THEN 1 ELSE 0 END), 0) AS total_status_inativo,
                    COALESCE(SUM(CASE WHEN ativo_no_relatorio = 'SIM' THEN quantidade_full ELSE 0 END), 0) AS total_full,
                    COALESCE(SUM(CASE WHEN ativo_no_relatorio = 'SIM' THEN quantidade_enviar_full ELSE 0 END), 0) AS total_enviar,
                    COALESCE(SUM(CASE WHEN ativo_no_relatorio = 'SIM' AND precisa_repor = 'SIM' THEN 1 ELSE 0 END), 0) AS qtd_repor,
                    COALESCE(SUM(CASE WHEN ativo_no_relatorio = 'SIM' AND estoque_minimo = 0 AND estoque_recomendado = 0 THEN 1 ELSE 0 END), 0) AS sem_config
                FROM anuncios_full
                """
            )
            metrics = cur.fetchone()
            cur.execute(
                """
                SELECT data_importacao_br, nome_arquivo, total_relatorio, qtd_novos, qtd_alterados, qtd_ausentes
                FROM importacoes_relatorios
                ORDER BY id DESC
                LIMIT 1
                """
            )
            ultima_importacao = cur.fetchone()

    return render_template(
        "index.html",
        rows=rows,
        metrics=metrics,
        busca=busca,
        filtro=filtro,
        ultima_importacao=ultima_importacao,
    )


@app.route("/produto/<int:produto_id>", methods=["GET", "POST"])
def produto(produto_id: int):
    init_db()
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT * FROM anuncios_full WHERE id = %s", (produto_id,))
            row = cur.fetchone()
            if row is None:
                flash("Produto não encontrado.", "danger")
                return redirect(url_for("index"))

            if request.method == "POST":
                estoque_minimo = int_safe(request.form.get("estoque_minimo"), 0)
                estoque_recomendado = int_safe(request.form.get("estoque_recomendado"), 0)
                observacao = request.form.get("observacao", "").strip()
                custo_produto = float_safe(request.form.get("custo_produto"), 0.0)
                custo_mercado_livre = float_safe(request.form.get("custo_mercado_livre"), 0.0)
                custo_impostos = float_safe(request.form.get("custo_impostos"), 0.0)
                cur.execute(
                    """
                    UPDATE anuncios_full
                    SET estoque_minimo = %s, estoque_recomendado = %s, observacao = %s,
                        custo_produto = %s, custo_mercado_livre = %s, custo_impostos = %s,
                        atualizado_em = CURRENT_TIMESTAMP
                    WHERE id = %s
                    """,
                    (estoque_minimo, estoque_recomendado, observacao, custo_produto,
                     custo_mercado_livre, custo_impostos, produto_id),
                )
                conn.commit()
                flash("Configuração do anúncio atualizada.", "success")
                return redirect(url_for("produto", produto_id=produto_id))

    return render_template("produto.html", row=row)


@app.route("/importar", methods=["GET", "POST"])
def importar():
    init_db()
    resultado = None
    if request.method == "POST":
        file = request.files.get("arquivo")

        if not file or not file.filename:
            flash("Selecione um arquivo Excel.", "warning")
            return redirect(url_for("importar"))

        filename = f"relatorio_{agora_brasil().strftime('%Y%m%d_%H%M%S')}.xlsx"
        path = UPLOAD_DIR / filename
        file.save(path)

        try:
            df_new = read_report(path)
            if df_new.empty:
                raise ValueError("Nenhum anúncio MLB encontrado na aba Anúncios.")

            alterados = []
            novos = []
            ausentes = []
            import_time = agora_brasil_str()

            with get_conn() as conn:
                with conn.cursor() as cur:
                    cur.execute("SELECT * FROM anuncios_full")
                    old_rows = cur.fetchall()
                    old_by_code = {r["codigo_anuncio"]: r for r in old_rows}
                    new_codes = set(df_new["codigo_anuncio"].tolist())
                    old_codes = set(old_by_code.keys())

                    cur.execute("UPDATE anuncios_full SET ativo_no_relatorio = 'NAO', atualizado_em = CURRENT_TIMESTAMP")

                    for _, r in df_new.iterrows():
                        codigo = r["codigo_anuncio"]
                        old = old_by_code.get(codigo)

                        if old:
                            imported_editable = {col: r.get(col) for col in EDITABLE_COLUMNS}
                            editable = merge_editable(old, imported_editable)

                            changes = {}
                            for col in [
                                "numero_produto", "titulo", "variacoes", "quantidade_full", "preco", "moeda",
                                "condicao", "forma_entrega", "tipo_anuncio", "status", "altura_cm", "largura_cm",
                                "profundidade_cm", "peso_kg", "ativo_no_relatorio",
                                "estoque_minimo", "estoque_recomendado", "observacao",
                                "custo_produto", "custo_mercado_livre", "custo_impostos"
                            ]:
                                old_val = old.get(col)
                                if col == "ativo_no_relatorio":
                                    new_val = "SIM"
                                elif col in EDITABLE_COLUMNS:
                                    new_val = editable[col]
                                else:
                                    new_val = r[col]

                                old_norm = normalize_for_compare(old_val, col)
                                new_norm = normalize_for_compare(new_val, col)

                                if old_norm != new_norm:
                                    changes[col] = {"antes": old_norm, "depois": new_norm}

                            if changes:
                                alterados.append({"codigo_anuncio": codigo, "titulo": r["titulo"], "changes": changes})

                            cur.execute(
                                """
                                UPDATE anuncios_full
                                SET numero_produto=%s, titulo=%s, variacoes=%s, quantidade_full=%s, preco=%s, moeda=%s,
                                    condicao=%s, forma_entrega=%s, tipo_anuncio=%s, status=%s, altura_cm=%s, largura_cm=%s,
                                    profundidade_cm=%s, peso_kg=%s, estoque_minimo=%s, estoque_recomendado=%s,
                                    observacao=%s, custo_produto=%s, custo_mercado_livre=%s, custo_impostos=%s,
                                    ativo_no_relatorio='SIM', data_ultima_importacao=%s,
                                    atualizado_em=CURRENT_TIMESTAMP
                                WHERE codigo_anuncio=%s
                                """,
                                (
                                    r["numero_produto"], r["titulo"], r["variacoes"], int(r["quantidade_full"]),
                                    float(r["preco"]), r["moeda"], r["condicao"], r["forma_entrega"], r["tipo_anuncio"],
                                    r["status"], int(r["altura_cm"]), int(r["largura_cm"]), int(r["profundidade_cm"]),
                                    float(r["peso_kg"]), editable["estoque_minimo"], editable["estoque_recomendado"],
                                    editable["observacao"], editable["custo_produto"], editable["custo_mercado_livre"],
                                    editable["custo_impostos"], import_time, codigo,
                                ),
                            )
                        else:
                            novos.append({
                                "codigo_anuncio": codigo,
                                "titulo": r["titulo"],
                                "quantidade_full": int(r["quantidade_full"]),
                            })
                            cur.execute(
                                """
                                INSERT INTO anuncios_full (
                                    codigo_anuncio, numero_produto, titulo, variacoes, quantidade_full, preco, moeda, condicao,
                                    forma_entrega, tipo_anuncio, status, altura_cm, largura_cm, profundidade_cm, peso_kg,
                                    estoque_minimo, estoque_recomendado, observacao,
                                    custo_produto, custo_mercado_livre, custo_impostos, ativo_no_relatorio,
                                    data_primeira_importacao, data_ultima_importacao
                                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'SIM', %s, %s)
                                """,
                                (
                                    codigo, r["numero_produto"], r["titulo"], r["variacoes"], int(r["quantidade_full"]),
                                    float(r["preco"]), r["moeda"], r["condicao"], r["forma_entrega"], r["tipo_anuncio"],
                                    r["status"], int(r["altura_cm"]), int(r["largura_cm"]), int(r["profundidade_cm"]),
                                    float(r["peso_kg"]),
                                    int(r.get("estoque_minimo") or 0), int(r.get("estoque_recomendado") or 0),
                                    r.get("observacao") or "", float(r.get("custo_produto") or 0),
                                    float(r.get("custo_mercado_livre") or 0), float(r.get("custo_impostos") or 0),
                                    import_time, import_time,
                                ),
                            )

                    missing_codes = sorted(old_codes - new_codes)
                    for code in missing_codes:
                        ausentes.append({
                            "codigo_anuncio": code,
                            "titulo": old_by_code[code].get("titulo", ""),
                            "quantidade_full": old_by_code[code].get("quantidade_full", 0),
                        })

                    cur.execute(
                        """
                        INSERT INTO importacoes_relatorios (
                            data_importacao, data_importacao_br, nome_arquivo, total_relatorio,
                            qtd_novos, qtd_alterados, qtd_ausentes
                        ) VALUES (CURRENT_TIMESTAMP, %s, %s, %s, %s, %s, %s)
                        """,
                        (import_time, file.filename, int(len(df_new)), len(novos), len(alterados), len(ausentes)),
                    )

                conn.commit()

            resultado = {
                "total_relatorio": int(len(df_new)),
                "alterados": alterados,
                "novos": novos,
                "ausentes": ausentes,
                "data_importacao": import_time,
                "nome_arquivo": file.filename,
            }
            flash("Relatório importado com sucesso. Os dados agora ficam salvos no PostgreSQL.", "success")

        except Exception as e:
            flash(f"Erro ao importar relatório: {e}", "danger")

    return render_template("importar.html", resultado=resultado)


@app.route("/exportar")
def exportar():
    init_db()
    out_path = BASE_DIR / "reposicao_full.csv"

    # Exportação robusta:
    # - não depende da coluna gerada quantidade_enviar_full;
    # - recalcula a reposição na hora;
    # - trata NULL como zero;
    # - aceita variações de ativo_no_relatorio, como SIM/sim/S/true/1;
    # - exporta tudo que tem quantidade a enviar > 0 ou reposição = SIM.
    sql = """
        WITH base AS (
            SELECT
                codigo_anuncio,
                numero_produto,
                titulo,
                COALESCE(quantidade_full, 0)::INTEGER AS quantidade_full,
                COALESCE(estoque_minimo, 0)::INTEGER AS estoque_minimo,
                COALESCE(estoque_recomendado, 0)::INTEGER AS estoque_recomendado,
                GREATEST(
                    (CASE
                        WHEN COALESCE(estoque_recomendado, 0) > 0
                            THEN COALESCE(estoque_recomendado, 0)
                        ELSE COALESCE(estoque_minimo, 0)
                     END) - COALESCE(quantidade_full, 0),
                    0
                )::INTEGER AS quantidade_enviar_full,
                CASE
                    WHEN COALESCE(estoque_minimo, 0) > 0
                     AND COALESCE(quantidade_full, 0) <= COALESCE(estoque_minimo, 0)
                    THEN 'SIM'
                    ELSE 'NAO'
                END AS precisa_repor,
                COALESCE(preco, 0) AS preco,
                COALESCE(status, '') AS status,
                COALESCE(forma_entrega, '') AS forma_entrega,
                COALESCE(observacao, '') AS observacao,
                COALESCE(custo_produto, 0) AS custo_produto,
                COALESCE(custo_mercado_livre, 0) AS custo_mercado_livre,
                COALESCE(custo_impostos, 0) AS custo_impostos,
                COALESCE(ativo_no_relatorio, 'SIM') AS no_relatorio
            FROM anuncios_full
            WHERE UPPER(TRIM(COALESCE(ativo_no_relatorio, 'SIM'))) IN ('SIM', 'S', 'YES', 'TRUE', '1')
        )
        SELECT
            codigo_anuncio,
            numero_produto,
            titulo,
            quantidade_full,
            estoque_minimo,
            estoque_recomendado,
            quantidade_enviar_full,
            precisa_repor,
            preco,
            status,
            forma_entrega,
            observacao,
            custo_produto,
            custo_mercado_livre,
            custo_impostos
        FROM base
        WHERE quantidade_enviar_full > 0
           OR precisa_repor = 'SIM'
        ORDER BY quantidade_enviar_full DESC, titulo ASC
    """

    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(sql)
            rows = cur.fetchall()

    colunas = [
        "codigo_anuncio", "numero_produto", "titulo", "quantidade_full",
        "estoque_minimo", "estoque_recomendado", "quantidade_enviar_full",
        "precisa_repor", "preco", "status", "forma_entrega", "observacao",
        "custo_produto", "custo_mercado_livre", "custo_impostos"
    ]
    df = pd.DataFrame(rows, columns=colunas)

    # Mesmo que não haja itens, o CSV sai com cabeçalho para facilitar diagnóstico.
    df.to_csv(out_path, index=False, sep=";", encoding="utf-8-sig")
    return send_file(out_path, as_attachment=True, download_name="reposicao_full.csv")



@app.route("/exportar-configuracao")
def exportar_configuracao():
    """Gera planilha com todos os anúncios e colunas editáveis para preenchimento em lote."""
    init_db()
    out_path = BASE_DIR / "configuracao_anuncios_full.xlsx"
    sql = """
        SELECT
            codigo_anuncio AS CODIGO_ANUNCIO,
            numero_produto AS SKU,
            titulo AS TITULO,
            ativo_no_relatorio AS NO_RELATORIO,
            quantidade_full AS QUANTIDADE_FULL,
            preco AS PRECO,
            status AS STATUS,
            estoque_minimo AS ESTOQUE_MINIMO,
            estoque_recomendado AS ESTOQUE_RECOMENDADO,
            COALESCE(observacao, '') AS OBSERVACAO,
            COALESCE(custo_produto, 0) AS CUSTO_PRODUTO,
            COALESCE(custo_mercado_livre, 0) AS CUSTO_MERCADO_LIVRE,
            COALESCE(custo_impostos, 0) AS CUSTO_IMPOSTOS
        FROM anuncios_full
        ORDER BY ativo_no_relatorio DESC, titulo ASC
    """
    # Não usar pd.read_sql_query aqui porque a conexão usa row_factory=dict_row (psycopg3).
    # Com read_sql_query, o pandas pode interpretar as chaves do dicionário como valores
    # e gerar uma planilha com os nomes das colunas repetidos em todas as linhas.
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(sql)
            rows = cur.fetchall()

    df = pd.DataFrame(rows)

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Configuração")
        ws = writer.book["Configuração"]
        ws.freeze_panes = "A2"
        widths = {
            "A": 18, "B": 16, "C": 55, "D": 14, "E": 16, "F": 12, "G": 14,
            "H": 18, "I": 22, "J": 40, "K": 16, "L": 22, "M": 18,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
        # Destaca as colunas editáveis.
        from openpyxl.styles import Font, PatternFill
        header_fill = PatternFill("solid", fgColor="1F2937")
        editable_fill = PatternFill("solid", fgColor="FFF2CC")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for row in ws.iter_rows(min_row=2, min_col=8, max_col=13):
            for cell in row:
                cell.fill = editable_fill

    return send_file(out_path, as_attachment=True, download_name="configuracao_anuncios_full.xlsx")


@app.route("/importar-configuracao", methods=["GET", "POST"])
def importar_configuracao():
    """Importa planilha preenchida com estoque mínimo/recomendado, observação e custos."""
    init_db()
    resultado = None
    if request.method == "POST":
        file = request.files.get("arquivo")
        if not file or not file.filename:
            flash("Selecione a planilha de configuração.", "warning")
            return redirect(url_for("importar_configuracao"))

        filename = f"configuracao_{agora_brasil().strftime('%Y%m%d_%H%M%S')}.xlsx"
        path = UPLOAD_DIR / filename
        file.save(path)

        try:
            df = pd.read_excel(path, sheet_name=0)
            if df.empty:
                raise ValueError("Planilha vazia.")
            codigo_col = find_column(df, ["CODIGO_ANUNCIO", "CODIGO ANUNCIO", "CÓDIGO ANÚNCIO", "ITEM_ID"])
            if not codigo_col:
                raise ValueError("Coluna CODIGO_ANUNCIO não encontrada.")

            atualizados = []
            ignorados = []
            sem_alteracao = 0

            with get_conn() as conn:
                with conn.cursor() as cur:
                    for idx in df.index:
                        codigo = str(df.at[idx, codigo_col] or "").strip()
                        if not codigo or codigo.lower() == "nan":
                            continue
                        cur.execute("SELECT * FROM anuncios_full WHERE codigo_anuncio = %s", (codigo,))
                        old = cur.fetchone()
                        if not old:
                            ignorados.append(codigo)
                            continue

                        imported = extract_editable_values(df, idx)
                        editable = merge_editable(old, imported)

                        mudou = False
                        for col in EDITABLE_COLUMNS:
                            if normalize_for_compare(old.get(col), col) != normalize_for_compare(editable[col], col):
                                mudou = True
                                break

                        if not mudou:
                            sem_alteracao += 1
                            continue

                        cur.execute(
                            """
                            UPDATE anuncios_full
                            SET estoque_minimo=%s,
                                estoque_recomendado=%s,
                                observacao=%s,
                                custo_produto=%s,
                                custo_mercado_livre=%s,
                                custo_impostos=%s,
                                atualizado_em=CURRENT_TIMESTAMP
                            WHERE codigo_anuncio=%s
                            """,
                            (
                                editable["estoque_minimo"], editable["estoque_recomendado"], editable["observacao"],
                                editable["custo_produto"], editable["custo_mercado_livre"], editable["custo_impostos"], codigo,
                            ),
                        )
                        atualizados.append(codigo)
                conn.commit()

            resultado = {
                "arquivo": file.filename,
                "atualizados": atualizados,
                "ignorados": ignorados,
                "sem_alteracao": sem_alteracao,
            }
            flash("Planilha de configuração importada com sucesso.", "success")
        except Exception as e:
            flash(f"Erro ao importar configuração: {e}", "danger")

    return render_template("importar_configuracao.html", resultado=resultado)


@app.route("/diagnostico-exportacao")
def diagnostico_exportacao():
    """Tela simples para conferir se o banco tem itens que deveriam sair no CSV."""
    init_db()
    sql = """
        SELECT
            codigo_anuncio,
            titulo,
            COALESCE(quantidade_full, 0)::INTEGER AS quantidade_full,
            COALESCE(estoque_minimo, 0)::INTEGER AS estoque_minimo,
            COALESCE(estoque_recomendado, 0)::INTEGER AS estoque_recomendado,
            GREATEST(
                (CASE
                    WHEN COALESCE(estoque_recomendado, 0) > 0 THEN COALESCE(estoque_recomendado, 0)
                    ELSE COALESCE(estoque_minimo, 0)
                 END) - COALESCE(quantidade_full, 0),
                0
            )::INTEGER AS quantidade_enviar_calculada,
            CASE
                WHEN COALESCE(estoque_minimo, 0) > 0
                 AND COALESCE(quantidade_full, 0) <= COALESCE(estoque_minimo, 0)
                THEN 'SIM'
                ELSE 'NAO'
            END AS precisa_repor_calculado,
            ativo_no_relatorio,
            status
        FROM anuncios_full
        ORDER BY quantidade_enviar_calculada DESC, titulo ASC
    """
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(sql)
            rows = cur.fetchall()
    return {"total_linhas": len(rows), "amostra": rows[:20]}


if __name__ == "__main__":
    init_db()
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=True)
